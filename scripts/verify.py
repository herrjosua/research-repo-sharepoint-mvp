"""Re-open every built file and check it. Exits non-zero on any failure."""
from __future__ import annotations

import csv
import re
import shutil
import subprocess
import sys

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

from common import FORMAT, INDEX_OUT_NAME, OUT, BUILD, load_records
from build_pptx import parse_slides

failures: list[str] = []


def check(cond: bool, msg: str) -> None:
    print(("  ok   " if cond else "  FAIL ") + msg)
    if not cond:
        failures.append(msg)


def verify_docx(rec) -> None:
    path = OUT / rec.out_name
    doc = Document(str(path))
    h1 = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
    h2 = [p.text for p in doc.paragraphs if p.style.name == "Heading 2"]
    src_h2 = re.findall(r"^## (.+)$", rec.body, flags=re.M)
    check(len(h1) == 1 and rec.id in h1[0], f"{rec.id}: one Heading 1 containing the ID")
    check(h2 == src_h2, f"{rec.id}: {len(src_h2)} Heading 2 sections match the Markdown")
    check(len(doc.tables) >= 1 and doc.tables[0].rows[0].cells[0].text.strip() == "Field",
          f"{rec.id}: header table present")
    check(doc.core_properties.title == rec.title, f"{rec.id}: document title property set")
    md_items = len(re.findall(r"^(?:- |\d+\. )", rec.body, flags=re.M))
    docx_items = sum(1 for p in doc.paragraphs if p._p.pPr is not None and p._p.pPr.numPr is not None)
    check(docx_items == md_items, f"{rec.id}: {md_items} list items kept as real list items (found {docx_items})")
    src_tables = len(re.findall(r"^\|[-| :]+\|$", rec.body, flags=re.M))
    check(len(doc.tables) == src_tables, f"{rec.id}: {src_tables} tables kept")


def verify_pdf(rec) -> None:
    path = OUT / rec.out_name
    check(path.exists() and path.stat().st_size > 0, f"{rec.id}: PDF exists")
    if shutil.which("pdftotext"):
        text = subprocess.run(["pdftotext", "-layout", str(path), "-"], capture_output=True, text=True).stdout
        check(rec.id in text and "cancel" in text.lower(), f"{rec.id}: PDF has a text layer with the content")
        check("If you cancel one week you cancel them all" in " ".join(text.split()), f"{rec.id}: key quote found in PDF text")
    else:
        print("  skip pdftotext not installed: text layer not checked")


def verify_usage(rec) -> None:
    wb = load_workbook(OUT / rec.out_name)
    check(wb.sheetnames == ["README", "Summary", "By size", "By room"], f"{rec.id}: sheets {wb.sheetnames}")
    rooms = wb["By room"]
    check(rooms.max_row == 9, f"{rec.id}: 8 room rows")
    check(len(rooms.tables) == 1 and len(wb["Summary"].tables) == 1, f"{rec.id}: Excel Tables defined")
    summary = {r[0]: r[1] for r in wb["Summary"].iter_rows(min_row=2, values_only=True)}
    md = rec.body
    check(summary["Unused hours"] == 109 and "109 of 352" in md, f"{rec.id}: unused hours 109 matches Markdown")
    check(summary["Unused share of booked hours"] == 31.0 and "31.0%" in md, f"{rec.id}: unused share 31.0% matches Markdown")
    check(summary["No-show share of bookings"] == 29.1 and "29.1%" in md, f"{rec.id}: no-show share 29.1% matches Markdown")
    sizes = {r[0]: r[5] for r in wb["By size"].iter_rows(min_row=2, values_only=True)}
    check(sizes["Small (4 seats)"] == 40.4 and "40.4%" in md, f"{rec.id}: small-room unused 40.4% matches Markdown")
    keel = [r for r in rooms.iter_rows(min_row=2, values_only=True) if r[1] == "Keel"][0]
    check(keel[8] == 44.1, f"{rec.id}: Keel unused 44.1%")
    # no formulas: values must be readable without recalculation
    check(not any(isinstance(c.value, str) and c.value.startswith("=") for ws in wb for row in ws.iter_rows() for c in row),
          f"{rec.id}: no formulas (values only)")


def verify_pptx(rec) -> None:
    prs = Presentation(str(OUT / rec.out_name))
    src = parse_slides(rec.body)
    check(len(prs.slides) == len(src), f"{rec.id}: {len(src)} slides")
    for slide, s in zip(prs.slides, src):
        title = slide.shapes.title.text if slide.shapes.title is not None else ""
        notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
        check(title == s["title"] and len(notes) > 40, f"{rec.id}: slide {s['num']} title and notes present")
        if s["table"]:
            tables = [sh for sh in slide.shapes if sh.has_table]
            check(len(tables) == 1 and len(tables[0].table.rows) == len(s["table"])
                  and tables[0]._element.nvGraphicFramePr.cNvPr.get("descr"),
                  f"{rec.id}: slide {s['num']} table with alt text")


def verify_index(records) -> None:
    wb = load_workbook(OUT / INDEX_OUT_NAME)
    ws = wb["Index"]
    check(ws.max_row == len(records) + 1, f"index: {len(records)} rows")
    files = {r[-1] for r in ws.iter_rows(min_row=2, values_only=True)}
    check(files == {r.out_name for r in records}, "index: File column matches delivered files")
    check(all(r[-2] for r in ws.iter_rows(min_row=2, values_only=True)), "index: every row has a summary")


def verify_metadata(records) -> None:
    with open(BUILD / "metadata.csv", newline="", encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
    names = {r["File name"] for r in rows}
    delivered = {p.name for p in OUT.iterdir()}
    check(names == delivered, "metadata.csv lists exactly the delivered files")
    f01 = [r for r in rows if r["ID"] == "F-01"][0]
    check(f01["Prevalence"] == "5 of 6 participants" and f01["Valence"] == "pain-point", "metadata.csv F-01 values")


def main() -> int:
    records = load_records()
    print(f"Verifying {len(records)} records + index in {OUT}")
    for rec in records:
        print(rec.id)
        fmt = FORMAT[rec.id]
        if fmt == "docx":
            verify_docx(rec)
        elif fmt == "pdf":
            verify_pdf(rec)
        elif fmt == "xlsx":
            verify_usage(rec)
        elif fmt == "pptx":
            verify_pptx(rec)
    print("index and metadata")
    verify_index(records)
    verify_metadata(records)
    expected = len(records) + 1
    check(len(list(OUT.iterdir())) == expected, f"{expected} files in build/sharepoint")
    print(f"\n{len(failures)} failure(s)")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
