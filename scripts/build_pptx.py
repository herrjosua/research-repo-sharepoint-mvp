"""Readout Markdown (R-01) -> PowerPoint.

Each '## Slide N: <title>' section becomes a slide. The title is a real title placeholder, the
'On the slide' content becomes bullets or a table, and 'Speaker notes' go in the notes pane.
"""
from __future__ import annotations

import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

from common import Record, parse_md_table

NAVY = RGBColor(0x1F, 0x2A, 0x44)
GREY = RGBColor(0x59, 0x59, 0x59)
FOOTER = "Fictional demo data  |  S-01 Meeting Room Booking Study  |  R-01"


def _clean(s: str) -> str:
    return re.sub(r"\*\*|`", "", s).strip()


def parse_slides(body: str) -> list[dict]:
    parts = re.split(r"^## Slide (\d+): (.+)$", body, flags=re.M)
    slides = []
    for i in range(1, len(parts), 3):
        num, title, content = int(parts[i]), parts[i + 1].strip(), parts[i + 2]
        on_slide, _, notes = content.partition("**Speaker notes:**")
        on_slide = on_slide.replace("**On the slide:**", "")
        table_lines = [l for l in on_slide.splitlines() if l.strip().startswith("|")]
        bullets: list[str] = []
        for line in on_slide.splitlines():
            if line.startswith("- "):
                bullets.append(line[2:].strip())
            elif line.startswith("  ") and bullets and not line.strip().startswith("|"):
                bullets[-1] += " " + line.strip()
        slides.append({
            "num": num,
            "title": _clean(title),
            "bullets": [_clean(b) for b in bullets],
            "table": [[_clean(c) for c in r] for r in parse_md_table(table_lines)] if table_lines else None,
            "notes": " ".join(_clean(l) for l in notes.strip().splitlines() if l.strip()),
        })
    return slides


def _place(shape, left, top, width, height):
    shape.left, shape.top, shape.width, shape.height = Inches(left), Inches(top), Inches(width), Inches(height)


def _style_title(shape, size=30):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for p in tf.paragraphs:
        p.alignment = PP_ALIGN.LEFT
        for r in p.runs:
            r.font.size = Pt(size)
            r.font.bold = True
            r.font.color.rgb = NAVY


def _bullets(tf, items, size):
    tf.word_wrap = True
    for i, text in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = text
        p.space_after = Pt(8)
        for r in p.runs:
            r.font.size = Pt(size)


def _footer(slide):
    tb = slide.shapes.add_textbox(Inches(0.6), Inches(7.0), Inches(12.1), Inches(0.3))
    tb.name = "Footer"
    p = tb.text_frame.paragraphs[0]
    p.text = FOOTER
    for r in p.runs:
        r.font.size = Pt(11)
        r.font.color.rgb = GREY


def _table(slide, rows, top):
    n_rows, n_cols = len(rows), len(rows[0])
    height = 0.6 * n_rows
    gf = slide.shapes.add_table(n_rows, n_cols, Inches(0.6), Inches(top), Inches(12.1), Inches(height))
    gf.name = "Data table"
    gf._element.nvGraphicFramePr.cNvPr.set(
        "descr", "Table with columns: " + ", ".join(rows[0]) + f". {n_rows - 1} data rows.")
    table = gf.table
    weights = [max(len(r[c]) for r in rows) + 6 for c in range(n_cols)]
    total = sum(weights)
    for c in range(n_cols):
        table.columns[c].width = Inches(12.1 * weights[c] / total)
    for ri, row in enumerate(rows):
        for ci, text in enumerate(row):
            cell = table.cell(ri, ci)
            cell.text = text
            for p in cell.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(18)
    return top + height


def build_pptx(rec: Record, out: Path) -> Path:
    slides = parse_slides(rec.body)
    if not slides:
        raise ValueError("no slides found in readout")
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)

    for s in slides:
        if s["num"] == 1:
            slide = prs.slides.add_slide(prs.slide_layouts[0])
            slide.shapes.title.text = s["title"]
            _place(slide.shapes.title, 0.8, 1.8, 11.7, 2.0)
            _style_title(slide.shapes.title, 36)
            sub = slide.placeholders[1]
            _place(sub, 0.8, 4.1, 11.7, 1.5)
            sub.text_frame.text = ""
            _bullets(sub.text_frame, s["bullets"], 24)
            for p in sub.text_frame.paragraphs:
                p.alignment = PP_ALIGN.LEFT
        elif s["table"]:
            slide = prs.slides.add_slide(prs.slide_layouts[5])  # Title Only
            slide.shapes.title.text = s["title"]
            _place(slide.shapes.title, 0.6, 0.4, 12.1, 1.3)
            _style_title(slide.shapes.title)
            bottom = _table(slide, s["table"], 1.9)
            if s["bullets"]:
                tb = slide.shapes.add_textbox(Inches(0.6), Inches(bottom + 0.3), Inches(12.1), Inches(1.5))
                tb.name = "Notes under table"
                _bullets(tb.text_frame, s["bullets"], 24)
        else:
            slide = prs.slides.add_slide(prs.slide_layouts[1])  # Title and Content
            slide.shapes.title.text = s["title"]
            _place(slide.shapes.title, 0.6, 0.4, 12.1, 1.3)
            _style_title(slide.shapes.title)
            body = slide.placeholders[1]
            _place(body, 0.6, 1.9, 12.1, 4.9)
            body.text_frame.text = ""
            _bullets(body.text_frame, s["bullets"], 24)
        slide.notes_slide.notes_text_frame.text = s["notes"]
        _footer(slide)

    cp = prs.core_properties
    cp.title = rec.title
    cp.author = str(rec.meta.get("author", ""))
    cp.subject = "S-01 Meeting Room Booking Study"
    cp.keywords = rec.id
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    return out
